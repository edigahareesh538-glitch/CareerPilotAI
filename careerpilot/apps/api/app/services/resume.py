import json
import uuid
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.resume import Resume, ResumeStatus, ResumeVersion
from app.models.user import User
from app.core.config import get_settings

settings = get_settings()


def extract_text_from_pdf(file_path: Path) -> str:
    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
    return "\n".join(text_parts)


def extract_text_from_docx(file_path: Path) -> str:
    doc = Document(file_path)
    text_parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_parts.append(cell.text)
    return "\n".join(text_parts)


def extract_text(file_path: Path, mime_type: str) -> str:
    if mime_type == "application/pdf":
        return extract_text_from_pdf(file_path)
    elif mime_type in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ):
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {mime_type}")


def parse_resume_text(text: str) -> dict[str, Any]:
    lines = [line.strip() for line in text.split("\n") if line.strip()]

    result = {
        "candidate": {
            "name": "",
            "email": "",
            "phone": "",
            "location": "",
            "linkedin": "",
            "github": "",
            "portfolio": "",
            "summary": "",
            "education": [],
            "experience": [],
            "skills": [],
            "projects": [],
            "certifications": [],
        },
        "target_roles": [],
        "technical_skills": [],
        "soft_skills": [],
        "years_experience": 0,
    }

    current_section = None
    for line in lines:
        lower = line.lower()
        if any(kw in lower for kw in ["education", "academic"]):
            current_section = "education"
            continue
        elif any(kw in lower for kw in ["experience", "employment", "work history"]):
            current_section = "experience"
            continue
        elif any(kw in lower for kw in ["skill", "technolog", "tool", "language"]):
            current_section = "skills"
            continue
        elif any(kw in lower for kw in ["project", "portfolio"]):
            current_section = "projects"
            continue
        elif any(kw in lower for kw in ["certification", "certificate", "license"]):
            current_section = "certifications"
            continue
        elif any(kw in lower for kw in ["summary", "objective", "profile"]):
            current_section = "summary"
            continue

        if current_section == "education":
            result["candidate"]["education"].append(line)
        elif current_section == "experience":
            result["candidate"]["experience"].append(line)
        elif current_section == "skills":
            skills = [s.strip() for s in line.replace(",", "\n").split("\n") if s.strip()]
            result["candidate"]["skills"].extend(skills)
        elif current_section == "projects":
            result["candidate"]["projects"].append(line)
        elif current_section == "certifications":
            result["candidate"]["certifications"].append(line)
        elif current_section == "summary":
            result["candidate"]["summary"] += " " + line
        elif "@" in line and not result["candidate"]["email"]:
            result["candidate"]["email"] = line
        elif any(keyword in lower for keyword in ["linkedin", "github"]):
            if "linkedin" in lower:
                result["candidate"]["linkedin"] = line
            elif "github" in lower:
                result["candidate"]["github"] = line

    if not result["candidate"]["name"] and lines:
        result["candidate"]["name"] = lines[0]

    result["candidate"]["skills"] = list(set(result["candidate"]["skills"]))
    result["technical_skills"] = result["candidate"]["skills"]
    result["soft_skills"] = []

    return result


async def process_resume(db: AsyncSession, resume_id: uuid.UUID) -> ResumeVersion:
    result = await db.execute(select(Resume).where(Resume.id == resume_id))
    resume = result.scalar_one_or_none()
    if not resume:
        raise ValueError("Resume not found")

    resume.status = ResumeStatus.PARSING
    await db.flush()

    try:
        from app.services.storage import download_file
        file_path = await download_file(resume.storage_key)

        text = extract_text(file_path, resume.mime_type)
        parsed_data = parse_resume_text(text)

        version_result = await db.execute(
            select(ResumeVersion).where(ResumeVersion.resume_id == resume_id).order_by(ResumeVersion.version.desc())
        )
        latest_version = version_result.scalars().first()
        next_version = (latest_version.version + 1) if latest_version else 1

        if latest_version:
            latest_version.is_active = False

        resume_version = ResumeVersion(
            resume_id=resume_id,
            parsed_json=parsed_data,
            version=next_version,
            is_active=True,
        )
        db.add(resume_version)

        resume.status = ResumeStatus.PARSED
        await db.flush()

        return resume_version

    except Exception as e:
        resume.status = ResumeStatus.FAILED
        resume.error_message = str(e)
        await db.flush()
        raise